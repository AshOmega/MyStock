import io
import json
import sqlite3
from datetime import date, datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pydantic import ValidationError

from app import main
from app.evidence import (
    ChartConfirmation, chart_assessment, chart_suggestions, company_key,
    financial_assessment, workbook_extract,
)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "DB_PATH", tmp_path / "evidence.sqlite3")
    with TestClient(main.app, base_url="http://localhost") as client:
        yield client


def workbook(company="Test Company Limited", change=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data Sheet"
    labels = {1: "COMPANY NAME", 16: "Report Date", 17: "Sales", 30: "Net profit",
              41: "Report Date", 42: "Sales", 49: "Net profit", 56: "Report Date",
              57: "Equity Share Capital", 58: "Reserves", 59: "Borrowings",
              81: "Report Date", 82: "Cash from Operating Activity",
              93: "Adjusted Equity Shares in Cr"}
    for row, label in labels.items():
        sheet.cell(row, 1, label)
    sheet["B1"] = company
    sheet["B8"] = 100
    sheet["B9"] = 10000
    year = date.today().year - 1
    for index, col in enumerate(range(2, 6)):
        for row in (16, 41, 56, 81):
            sheet.cell(row, col, datetime(year - 3 + index, 3, 31))
        for row, value in {17: 100 * 1.1 ** index, 30: 20 * 1.1 ** index,
                           42: 40, 49: 8, 57: 10, 58: 90, 59: 50,
                           82: 30, 93: 10}.items():
            sheet.cell(row, col, value)
    if change:
        change(sheet)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def preview_workbook(client, content=None):
    return client.post("/api/evidence/preview/workbook",
                       files={"file": ("file with % & ; / unicode.xlsx", content or workbook(),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def confirm_workbook(client, evidence_id, **changes):
    values = {"symbol": "TEST", "basis": "consolidated", "units": "INR_crore",
              "company_type": "non_financial", "identity_confirmed": True}
    values.update(changes)
    return client.post(f"/api/evidence/{evidence_id}/workbook", json=values)


def chart_values(**changes):
    values = {"symbol": "TEST", "company": "Test Company LTD", "timeframe": "1W",
              "captured_at": datetime.now(timezone.utc).isoformat(), "price": 100,
              "volume": 1000, "sma20": 120, "sma50": 130, "sma200": 150,
              "bar_complete": False, "identity_confirmed": True}
    values.update(changes)
    return values


def preview_chart(client, monkeypatch):
    image = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(image, format="PNG")
    monkeypatch.setattr(main, "chart_extract", lambda content: {
        "ocr_text": "Synthetic chart", "suggestions": {}, "warnings": ["Confirm labels."]})
    return client.post("/api/evidence/preview/chart", files={"file": ("test.png", image.getvalue(), "image/png")})


def test_workbook_metrics_units_provenance():
    payload = workbook_extract(workbook())
    result = financial_assessment(payload, "non_financial")
    assert result["score"] == 6
    assert result["status"] == "fundamentals_available"
    assert result["annual_age_days"] > 180
    assert result["metrics"][0]["value"] == pytest.approx(10)
    assert result["metrics"][1]["value"] == pytest.approx(10)
    assert result["metrics"][4]["value"] == 300_000_000
    assert result["metrics"][0]["cells"] == ["Data Sheet!B17", "Data Sheet!E17"]
    assert payload["undated_price"] == 100
    assert financial_assessment(payload, "bank")["status"] == "unsupported_sector"


@pytest.mark.parametrize("ref,value", [
    ("E17", "=1+1"), ("E30", None), ("E58", None), ("E82", None),
    ("B17", -100), ("B93", 0), ("E93", 0),
    ("E56", datetime(2000, 3, 31)), ("E81", datetime(2000, 3, 31)),
])
def test_formula_missing_negative_and_misaligned(ref, value):
    payload = workbook_extract(workbook(change=lambda s: setattr(s[ref], "value", value)))
    result = financial_assessment(payload, "non_financial")
    assert result["status"] == "partial_fundamentals"
    assert any(m["passed"] is None for m in result["metrics"])


def test_nonpositive_profit_and_zero_equity():
    payload = workbook_extract(workbook(change=lambda s: setattr(s["E30"], "value", -1)))
    result = financial_assessment(payload, "non_financial")
    assert next(m for m in result["metrics"] if m["key"] == "cash_ratio")["passed"] is False
    payload = workbook_extract(workbook(change=lambda s: setattr(s["E58"], "value", -10)))
    assert financial_assessment(payload, "non_financial")["status"] == "partial_fundamentals"


@pytest.mark.parametrize("ref,value", [
    ("A17", "Unknown sales"), ("E16", datetime(2099, 1, 1)),
    ("E16", "bad date"), ("E16", 0), ("E16", 1e308),
    ("C16", datetime(date.today().year - 4, 3, 31)),
    ("C41", datetime(date.today().year - 4, 3, 31)),
])
def test_unsupported_and_invalid_workbook(ref, value):
    with pytest.raises(ValueError):
        workbook_extract(workbook(change=lambda s: setattr(s[ref], "value", value)))


def test_empty_and_corrupt_uploads(client):
    assert preview_workbook(client, b"bad file").status_code == 422
    assert client.post("/api/evidence/preview/other", files={"file": ("x", b"x")}).status_code == 422
    assert client.post("/api/evidence/preview/chart", files={"file": ("x.png", b"bad")}).status_code == 422
    assert client.post("/api/evidence/preview/workbook",
                       files={"file": ("x.xlsx", b"x" * 10_000_001)}).status_code == 422


def test_confirmation_identity_and_persistence(client, monkeypatch):
    preview = preview_workbook(client)
    assert preview.status_code == 200
    evidence_id = preview.json()["id"]
    assert client.get("/api/evidence").json() == []
    assert confirm_workbook(client, evidence_id).status_code == 200
    assert confirm_workbook(client, evidence_id).status_code == 409
    assert confirm_workbook(client, 999).status_code == 404
    chart_id = preview_chart(client, monkeypatch).json()["id"]
    assert client.get(f"/api/evidence/{chart_id}/image").headers["content-type"] == "image/png"
    assert client.get("/api/evidence/999/image").status_code == 404
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values()).status_code == 200
    data = client.get("/api/evidence").json()
    assert len(data) == 1
    assert data[0]["symbol"] == "TEST"
    assert data[0]["workbook"]["assessment"]["score"] == 6
    assert data[0]["chart"]["assessment"]["status"] == "below_all_averages"
    assert data[0]["entry_status"] == "not_established"
    assert client.get("/api/stocks").json() == []
    assert client.get("/api/portfolio").json()["cash"] == 300000
    with main.db() as connection:
        rows = connection.execute("SELECT kind,symbol,content FROM evidence WHERE confirmed_at IS NOT NULL").fetchall()
        assert all(r["symbol"] == "TEST" and len(r["content"]) > 0 for r in rows)
    with TestClient(main.app, base_url="http://localhost") as restarted:
        assert restarted.get("/api/evidence").json()[0]["chart"]["id"] == chart_id


@pytest.mark.parametrize("changes", [
    {"symbol": "lowercase"}, {"symbol": "../TCS"}, {"symbol": "TCS; DROP TABLE evidence"},
    {"units": "INR"}, {"basis": "unknown"}, {"identity_confirmed": False},
    {"identity_confirmed": None}, {"company_type": "unknown"},
])
def test_invalid_workbook_confirmation(client, changes):
    evidence_id = preview_workbook(client).json()["id"]
    assert confirm_workbook(client, evidence_id, **changes).status_code == 422
    assert client.get("/api/evidence").json() == []


def test_collision_guards_and_history(client):
    evidence_id = preview_workbook(client).json()["id"]
    assert confirm_workbook(client, evidence_id).status_code == 200
    other = preview_workbook(client, workbook("Another Company")).json()["id"]
    assert confirm_workbook(client, other).status_code == 422
    same = preview_workbook(client).json()["id"]
    assert confirm_workbook(client, same, symbol="OTHER").status_code == 422
    assert confirm_workbook(client, same, basis="standalone").status_code == 422
    assert confirm_workbook(client, same, company_type="bank").status_code == 422
    assert confirm_workbook(client, same).status_code == 200
    older = preview_workbook(client, workbook(change=lambda s: setattr(
        s["E16"], "value", datetime(date.today().year - 2, 3, 30)))).json()["id"]
    assert confirm_workbook(client, older).status_code == 422
    with main.db() as connection:
        assert connection.execute("SELECT COUNT(*) FROM evidence WHERE symbol='TEST'").fetchone()[0] == 2


def test_chart_must_correlate_with_workbook(client, monkeypatch):
    chart_id = preview_chart(client, monkeypatch).json()["id"]
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values()).status_code == 422
    evidence_id = preview_workbook(client).json()["id"]
    assert confirm_workbook(client, evidence_id).status_code == 200
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values(company="Wrong Ltd")).status_code == 422
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values()).status_code == 200
    another = preview_chart(client, monkeypatch).json()["id"]
    assert client.post(f"/api/evidence/{another}/chart", json=chart_values(
        captured_at=(datetime.now(timezone.utc) - timedelta(days=10)).isoformat())).status_code == 422
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values()).status_code == 409


@pytest.mark.parametrize("changes", [
    {"timeframe": "1H"}, {"price": 0}, {"sma20": -1}, {"volume": -1},
    {"identity_confirmed": False}, {"captured_at": "2026-01-01T10:00:00"},
    {"captured_at": "2099-01-01T10:00:00+05:30"}, {"price": float("nan")},
])
def test_chart_validation(changes):
    with pytest.raises(ValidationError):
        ChartConfirmation(**chart_values(**changes))


@pytest.mark.parametrize("values,status", [
    ({"sma20": None}, "insufficient_chart_data"),
    ({"price": 120}, "below_all_averages"),
    ({"price": 125}, "mixed_trend"),
    ({"price": 160, "bar_complete": True}, "above_all_averages"),
])
def test_chart_trend_context(values, status):
    assessment = chart_assessment(chart_values(**values))
    assert assessment["status"] == status
    assert ("completed bar" if values.get("bar_complete") else "incomplete") in assessment["warning"]


def test_ocr_parser_preserves_timeframe_and_units():
    text = """created with TradingView.com, Oct 01, 2026 15:11 UTC+5:30
Tata Consultancy Services Limited • 1W • NSE O2,100.0 H2,101.0 L2,032.4 C2,076.2
Vol 13.3 M
SMAs (20, close, 50, close, 200, close) 2,215.4 2,547.5 3,167.0"""
    parsed = chart_suggestions(text)
    assert parsed == {"company": "Tata Consultancy Services Limited", "timeframe": "1W",
                      "captured_at": "2026-10-01T15:11:00+05:30", "price": 2076.2,
                      "volume": 13300000, "sma20": 2215.4, "sma50": 2547.5, "sma200": 3167}
    assert chart_suggestions("Unreadable chart") == {}
    assert "sma20" not in chart_suggestions("SMA (10, close) 2,215.4")
    assert company_key("TATA CONSULTANCY SERVICES LTD") == company_key("Tata Consultancy Services Limited")
    assert company_key("भारत Company LTD") != company_key("தமிழ் Company LTD")


def test_source_text_roundtrip(client):
    name = "Test % & ; / भारत Company Ltd"
    evidence_id = preview_workbook(client, workbook(name)).json()["id"]
    assert confirm_workbook(client, evidence_id).status_code == 200
    assert client.get("/api/evidence").json()[0]["company"] == name
    assert client.get("/api/health").status_code == 200


def test_latest_is_confirmation_order_not_upload_order(client):
    first = preview_workbook(client, workbook(change=lambda s: setattr(s["E17"], "value", 250))).json()["id"]
    second = preview_workbook(client).json()["id"]
    assert confirm_workbook(client, second).status_code == 200
    assert confirm_workbook(client, first).status_code == 200
    current = client.get("/api/evidence").json()[0]["workbook"]
    assert current["id"] == first
    assert current["payload"]["annual"][-1]["values"]["sales"] == 250


def test_company_linking_without_symbol(client, monkeypatch):
    preview = preview_workbook(client).json()
    confirmation = {"basis": "consolidated", "units": "INR_crore", "company_type": "non_financial",
                    "identity_confirmed": True}
    response = client.post(f"/api/evidence/{preview['id']}/workbook", json=confirmation)
    assert response.status_code == 200
    company_id = response.json()["company_id"]
    identity = client.get("/api/evidence").json()[0]
    assert identity["symbol"] is None
    assert identity["id"] == company_id
    monkeypatch.setattr(main, "chart_extract", lambda content: {
        "ocr_text": "Test Company LTD", "suggestions": {"company": "TEST COMPANY LTD"}, "warnings": []})
    image = io.BytesIO()
    Image.new("RGB", (10, 10)).save(image, format="PNG")
    chart = client.post("/api/evidence/preview/chart", files={"file": ("chart.png", image.getvalue())}).json()
    assert chart["suggested_company_id"] == company_id
    values = chart_values(company_id=company_id)
    del values["symbol"]
    assert client.post(f"/api/evidence/{chart['id']}/chart", json=values).status_code == 200
    identity = client.get("/api/evidence").json()[0]
    assert identity["chart"]["id"] == chart["id"]
    assert identity["symbol"] is None
    assert client.get("/api/portfolio").json()["cash"] == 300000
    with main.db() as connection:
        linked = connection.execute("SELECT company_id FROM evidence WHERE confirmed_at IS NOT NULL").fetchall()
        assert len(linked) == 2 and all(row["company_id"] == company_id for row in linked)


def test_chart_preview_does_not_guess_a_company(client, monkeypatch):
    assert confirm_workbook(client, preview_workbook(client).json()["id"]).status_code == 200
    monkeypatch.setattr(main, "chart_extract", lambda content: {
        "ocr_text": "Another Company LTD", "suggestions": {"company": "Another Company LTD"}, "warnings": []})
    chart = client.post("/api/evidence/preview/chart", files={"file": ("TEST.png", b"synthetic")}).json()
    assert chart["suggested_company_id"] is None
    monkeypatch.setattr(main, "chart_extract", lambda content: {
        "ocr_text": "", "suggestions": {}, "warnings": ["Unreadable"]})
    chart = client.post("/api/evidence/preview/chart", files={"file": ("TEST.png", b"synthetic")}).json()
    assert chart["suggested_company_id"] is None


def test_legacy_symbol_records_migrate_without_loss(client):
    payload = workbook_extract(workbook())
    old_workbook = {"symbol": "TEST", "basis": "consolidated", "units": "INR_crore",
                    "company_type": "non_financial", "identity_confirmed": True}
    with sqlite3.connect(main.DB_PATH) as connection:
        connection.executescript("""
            CREATE TABLE equities(symbol TEXT PRIMARY KEY,company TEXT NOT NULL,company_key TEXT NOT NULL UNIQUE);
            CREATE TABLE evidence(id INTEGER PRIMARY KEY,kind TEXT NOT NULL,symbol TEXT,filename TEXT NOT NULL,
                content BLOB NOT NULL,payload TEXT NOT NULL,confirmation TEXT,
                imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,confirmed_at TEXT);
        """)
        connection.execute("INSERT INTO equities VALUES (?,?,?)", ("TEST", payload["company"], company_key(payload["company"])))
        connection.execute("""INSERT INTO evidence(kind,symbol,filename,content,payload,confirmation,confirmed_at)
            VALUES ('workbook','TEST','old.xlsx',?,?,?,CURRENT_TIMESTAMP)""",
                           (b"original file bytes", json.dumps(payload), json.dumps(old_workbook)))
        connection.execute("""INSERT INTO evidence(kind,symbol,filename,content,payload,confirmation,confirmed_at)
            VALUES ('chart','TEST','old.png',?,?,?,CURRENT_TIMESTAMP)""",
                           (b"original image bytes", json.dumps({"ocr_text": "old chart"}),
                            json.dumps(chart_values())))
    result = client.get("/api/evidence")
    assert result.status_code == 200
    item = result.json()[0]
    assert item["symbol"] == "TEST"
    assert item["workbook"]["id"] == 1
    assert item["chart"]["id"] == 2
    assert client.get("/api/evidence").json()[0]["id"] == item["id"]
    with main.db() as connection:
        assert connection.execute("SELECT content FROM evidence WHERE id=1").fetchone()[0] == b"original file bytes"
        assert connection.execute("SELECT COUNT(*) FROM companies").fetchone()[0] == 1


def test_chart_company_and_legacy_symbol_cannot_disagree(client, monkeypatch):
    response = confirm_workbook(client, preview_workbook(client).json()["id"])
    company_id = response.json()["company_id"]
    chart_id = preview_chart(client, monkeypatch).json()["id"]
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values(
        company_id=company_id, symbol="OTHER")).status_code == 422
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values(
        company_id=999, symbol=None)).status_code == 422
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values(
        company_id=None, symbol=None)).status_code == 422


def test_drafts_resume_latest_files_without_implying_confirmation(client, monkeypatch):
    assert client.get("/api/evidence/drafts").json() == []
    first = preview_workbook(client).json()["id"]
    second = preview_workbook(client).json()["id"]
    chart_id = preview_chart(client, monkeypatch).json()["id"]
    drafts = client.get("/api/evidence/drafts").json()
    assert [d["id"] for d in drafts] == [second, chart_id]
    assert drafts[0]["payload"]["company"] == "Test Company Limited"
    assert client.get("/api/evidence").json() == []
    assert confirm_workbook(client, second).status_code == 200
    assert [d["id"] for d in client.get("/api/evidence/drafts").json()] == [chart_id]
    assert client.post(f"/api/evidence/{chart_id}/chart", json=chart_values()).status_code == 200
    assert client.get("/api/evidence/drafts").json() == []
    with main.db() as connection:
        assert connection.execute("SELECT confirmed_at FROM evidence WHERE id=?", (first,)).fetchone()[0] is None
