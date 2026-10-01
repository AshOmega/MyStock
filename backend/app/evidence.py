import io
import json
import math
import re
import subprocess
import tempfile
from xml.etree.ElementTree import ParseError
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Literal, Optional
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from PIL import Image, UnidentifiedImageError
from pydantic import Field, model_validator

from .models import InputModel


SYMBOL_PATTERN = r"^[A-Z0-9&._-]+$"


def company_key(value):
    words = re.findall(r"[^\W_]+", value.upper())
    return " ".join("LIMITED" if w == "LTD" else w for w in words)


class WorkbookConfirmation(InputModel):
    symbol: Optional[str] = Field(default=None, min_length=1, max_length=40, pattern=SYMBOL_PATTERN)
    basis: Literal["consolidated", "standalone"]
    units: Literal["INR_crore"]
    company_type: Literal["non_financial", "bank", "nbfc", "insurer"]
    identity_confirmed: Literal[True]


class ChartConfirmation(InputModel):
    company_id: Optional[int] = Field(default=None, gt=0)
    symbol: Optional[str] = Field(default=None, min_length=1, max_length=40, pattern=SYMBOL_PATTERN)
    company: str = Field(min_length=1, max_length=200)
    timeframe: Literal["1W", "1D"]
    captured_at: datetime
    price: float = Field(gt=0)
    volume: Optional[float] = Field(default=None, ge=0)
    sma20: Optional[float] = Field(default=None, gt=0)
    sma50: Optional[float] = Field(default=None, gt=0)
    sma200: Optional[float] = Field(default=None, gt=0)
    bar_complete: bool
    identity_confirmed: Literal[True]

    @model_validator(mode="after")
    def timestamp(self):
        if self.company_id is None and self.symbol is None:
            raise ValueError("Choose a confirmed workbook company")
        if self.captured_at.tzinfo is None:
            raise ValueError("Screenshot timestamp requires an explicit timezone offset")
        if self.captured_at > datetime.now(timezone.utc):
            raise ValueError("Screenshot capture cannot be in the future")
        return self


def number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value) if math.isfinite(value) else None


def workbook_extract(content):
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20_000_000:
                raise ValueError("Expanded XLSX exceeds 20 MB.")
            if any("vbaProject" in item.filename for item in archive.infolist()):
                raise ValueError("Macro-enabled workbooks are not accepted.")
        book = load_workbook(io.BytesIO(content), data_only=False, read_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError, ParseError) as error:
        raise ValueError("Invalid or unsupported XLSX archive.") from error
    try:
        if "Data Sheet" not in book.sheetnames:
            raise ValueError("Expected the Screener 'Data Sheet'; other XLSX layouts are unsupported.")
        sheet = book["Data Sheet"]
        if sheet.max_row > 200 or sheet.max_column > 40:
            raise ValueError("Unexpected Screener sheet dimensions; review the workbook layout.")
        cells = {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}
        labels = {
            "A1": "COMPANY NAME", "A16": "Report Date", "A17": "Sales",
            "A41": "Report Date", "A42": "Sales", "A49": "Net profit",
            "A30": "Net profit", "A56": "Report Date", "A57": "Equity Share Capital",
            "A58": "Reserves", "A59": "Borrowings", "A81": "Report Date",
            "A82": "Cash from Operating Activity", "A93": "Adjusted Equity Shares in Cr",
        }
        for ref, expected in labels.items():
            if cells.get(ref) != expected:
                raise ValueError(f"Unsupported workbook layout at {ref}: expected {expected}.")
        name = cells.get("B1")
        if not isinstance(name, str) or not name.strip() or name.startswith("="):
            raise ValueError("Company name is missing.")
        fields = {
            "sales": 17, "net_profit": 30, "capital": 57, "reserves": 58,
            "borrowings": 59, "operating_cash_flow": 82, "adjusted_shares_crore": 93,
        }
        annual = []
        quarterly = []
        warnings = []

        def period(ref):
            raw = cells.get(ref)
            if isinstance(raw, datetime):
                result = raw.date()
            elif isinstance(raw, date):
                result = raw
            elif number(raw) is not None:
                try:
                    converted = from_excel(raw, book.epoch)
                except (OverflowError, ValueError) as error:
                    raise ValueError(f"Invalid Excel reporting date in {ref}.") from error
                if not isinstance(converted, datetime):
                    raise ValueError(f"Invalid Excel reporting date in {ref}.")
                result = converted.date()
            elif raw is None:
                return None
            else:
                raise ValueError(f"Invalid reporting date in {ref}.")
            if result > date.today():
                raise ValueError(f"Future reporting period in {ref}.")
            return result

        for col in range(2, sheet.max_column + 1):
            letter = get_column_letter(col)
            end = period(f"{letter}16")
            if end:
                item = {"period": str(end), "cells": {}, "values": {}}
                for key, row in fields.items():
                    ref = f"{letter}{row}"
                    item["values"][key] = number(cells.get(ref))
                    item["cells"][key] = f"Data Sheet!{ref}"
                item["aligned"] = period(f"{letter}56") == end and period(f"{letter}81") == end
                annual.append(item)
            end = period(f"{letter}41")
            if end:
                quarterly.append({"period": str(end), "sales": number(cells.get(f"{letter}42")),
                                  "net_profit": number(cells.get(f"{letter}49")),
                                  "cells": [f"Data Sheet!{letter}42", f"Data Sheet!{letter}49"]})
        annual.sort(key=lambda x: x["period"])
        quarterly.sort(key=lambda x: x["period"])
        if len({a["period"] for a in annual}) != len(annual):
            raise ValueError("Duplicate annual reporting dates.")
        if len({q["period"] for q in quarterly}) != len(quarterly):
            raise ValueError("Duplicate quarterly reporting dates.")
        if not annual:
            raise ValueError("No annual financial observations found.")
        if not all(a["aligned"] for a in annual):
            warnings.append("Some statement dates are not aligned; affected ratios are unavailable.")
        warnings += ["Workbook quote has no observation date; it is not a current executable price.",
                     "Publication dates, governance, pledges, daily liquidity and fair value remain unknown.",
                     "EPS uses template adjusted shares, not verified diluted EPS. ROE uses closing equity.",
                     "Annual and quarterly periods are shown separately; no blanket age rule determines suitability."]
        return {
            "company": name.strip(), "sheets": book.sheetnames,
            "annual": annual, "quarterly": quarterly,
            "undated_price": number(cells.get("B8")),
            "market_cap_crore": number(cells.get("B9")),
            "warnings": warnings,
            "source": "User-uploaded Screener-format XLSX; attribution and basis require confirmation.",
        }
    finally:
        book.close()


def financial_assessment(payload, company_type):
    if company_type != "non_financial":
        return {"status": "unsupported_sector", "score": None, "metrics": [],
                "warning": "Financial-sector rules are not approved."}
    annual = payload["annual"]
    latest = annual[-1]
    v = latest["values"]
    metrics = []

    def add(key, title, value, passed, refs, method):
        if value is not None and not math.isfinite(value):
            value, passed = None, None
        metrics.append({"key": key, "title": title, "value": value, "passed": passed,
                        "cells": refs, "method": method})

    end = date.fromisoformat(latest["period"])
    start = next((a for a in annual if date.fromisoformat(a["period"]).year == end.year - 3
                  and a["period"][5:] == latest["period"][5:]), None)
    for key, title in (("sales", "3-year revenue CAGR (%)"), ("eps", "3-year template EPS CAGR (%)")):
        value = None
        refs = []
        if start:
            a, b = start["values"], v
            if key == "sales":
                x, y = a["sales"], b["sales"]
                refs = [start["cells"]["sales"], latest["cells"]["sales"]]
            else:
                x = a["net_profit"] / a["adjusted_shares_crore"] if a["net_profit"] is not None and a["adjusted_shares_crore"] and a["adjusted_shares_crore"] > 0 else None
                y = b["net_profit"] / b["adjusted_shares_crore"] if b["net_profit"] is not None and b["adjusted_shares_crore"] and b["adjusted_shares_crore"] > 0 else None
                refs = [start["cells"]["net_profit"], start["cells"]["adjusted_shares_crore"],
                        latest["cells"]["net_profit"], latest["cells"]["adjusted_shares_crore"]]
            if x is not None and y is not None and x > 0 and y > 0:
                value = ((y / x) ** (1 / 3) - 1) * 100
        add(key, title, value, value > 0 if value is not None else None, refs,
            "Positive-base annual endpoints three years apart; (end/start)^(1/3)-1.")
    equity = v["capital"] + v["reserves"] if v["capital"] is not None and v["reserves"] is not None and latest["aligned"] else None
    roe = v["net_profit"] / equity * 100 if equity and equity > 0 and v["net_profit"] is not None else None
    debt = v["borrowings"] / equity if equity and equity > 0 and v["borrowings"] is not None and v["borrowings"] >= 0 else None
    cash = v["operating_cash_flow"] if latest["aligned"] else None
    profit = v["net_profit"]
    ratio = cash / profit if cash is not None and profit is not None and profit > 0 else None
    add("roe", "Closing-equity ROE (%)", roe, roe >= 15 if roe is not None else None,
        [latest["cells"][k] for k in ("net_profit", "capital", "reserves")], "Annual profit / closing (capital + reserves) * 100.")
    add("debt", "Debt/equity", debt, debt <= 1 if debt is not None else None,
        [latest["cells"][k] for k in ("borrowings", "capital", "reserves")], "Borrowings / (capital + reserves).")
    add("cash", "Operating cash flow (INR)", cash * 10_000_000 if cash is not None else None,
        cash > 0 if cash is not None else None, [latest["cells"]["operating_cash_flow"]], "Matched annual cash flow; crore converted to INR.")
    add("cash_ratio", "Operating cash flow/net profit", ratio, ratio >= 0.8 if ratio is not None else (False if cash is not None and profit is not None else None),
        [latest["cells"]["operating_cash_flow"], latest["cells"]["net_profit"]], "Same annual period; positive profit required.")
    complete = all(m["passed"] is not None for m in metrics)
    return {"status": "fundamentals_available" if complete else "partial_fundamentals",
            "score": sum(m["passed"] is True for m in metrics), "metrics": metrics,
            "annual_period": latest["period"], "annual_age_days": (date.today() - end).days,
            "quarterly_period": payload["quarterly"][-1]["period"] if payload["quarterly"] else None,
            "quarterly_age_days": (date.today() - date.fromisoformat(payload["quarterly"][-1]["period"])).days if payload["quarterly"] else None,
            "warning": "Descriptive, unvalidated quality checklist. Age is disclosed, not a buy/sell gate. Annual cash flow is not mixed with quarterly profit."}


def chart_suggestions(text):
    output = {}
    company = re.search(r"([^\n]+?)\s*[·•]\s*(1[WDwDd])\s*[·•]\s*NSE", text)
    if company:
        output.update(company=company.group(1).strip(), timeframe=company.group(2).upper())
    capture = re.search(r"([A-Z][a-z]{2}\s+\d{1,2},?\s+\d{4})\s+(\d{2}:\d{2})\s+UTC([+-])(\d{1,2}):(\d{2})", text)
    if capture:
        output["captured_at"] = datetime.strptime(capture.group(1).replace(",", ""), "%b %d %Y").strftime("%Y-%m-%d") + "T" + capture.group(2) + ":00" + capture.group(3) + capture.group(4).zfill(2) + ":" + capture.group(5)
    close = re.search(r"\bC\s*([0-9][0-9,]*\.\d+)", text)
    if close:
        output["price"] = float(close.group(1).replace(",", ""))
    volume = re.search(r"\bVol\s*([0-9.]+)\s*([KMB])?", text)
    if volume:
        output["volume"] = float(volume.group(1)) * {"K": 1000, "M": 1_000_000, "B": 1_000_000_000, None: 1}[volume.group(2)]
    sma = re.search(r"SMAs?\s*\(20,\s*close,\s*50,\s*close,\s*200,\s*close\)\s*([0-9,]+\.\d+)\s+([0-9,]+\.\d+)\s+([0-9,]+\.\d+)", text)
    if sma:
        output.update({key: float(value.replace(",", "")) for key, value in zip(("sma20", "sma50", "sma200"), sma.groups())})
    return output


def chart_extract(content):
    try:
        with Image.open(io.BytesIO(content)) as image:
            if image.format not in ("PNG", "JPEG") or image.width * image.height > 20_000_000:
                raise ValueError("Use PNG/JPEG up to 20 million pixels.")
            image.verify()
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as error:
        raise ValueError("Invalid or oversized chart image.") from error
    with tempfile.NamedTemporaryFile(suffix=".png") as image_file:
        image_file.write(content)
        image_file.flush()
        try:
            result = subprocess.run(["swift", str(Path(__file__).with_name("chart_ocr.swift")), image_file.name],
                                    capture_output=True, text=True, timeout=90, check=False)
        except (OSError, subprocess.TimeoutExpired) as error:
            raise ValueError("Local Apple Vision OCR failed. Ensure Xcode Command Line Tools are available.") from error
    if result.returncode:
        raise ValueError("Local Apple Vision OCR failed: " + result.stderr[-1000:])
    rows = json.loads(result.stdout)
    rows.sort(key=lambda r: (-round(r["y"], 2), r["x"]))
    text = "\n".join(r["text"] for r in rows)
    return {"ocr_text": text, "ocr_lines": rows, "suggestions": chart_suggestions(text),
            "warnings": ["OCR is fallible: confirm every extracted label against the screenshot.",
                         "Capture time is not proof of quote freshness, bar date or bar completion.",
                         "Volume belongs to the displayed bar, not average daily traded value.",
                         "No candle-history, support/resistance or custom indicator is inferred from pixels."]}


def chart_assessment(values):
    reasons = []
    available = 0
    above = 0
    for key in ("sma20", "sma50", "sma200"):
        average = values.get(key)
        if average is not None:
            available += 1
            higher = values["price"] > average
            above += int(higher)
            reasons.append(f"Price {'above' if higher else 'at or below'} {key[3:]}-{values['timeframe']} SMA ({average:,.2f}).")
    status = "insufficient_chart_data" if available < 3 else (
        "below_all_averages" if above == 0 else "above_all_averages" if above == 3 else "mixed_trend")
    capture_age = (datetime.now(timezone.utc) - datetime.fromisoformat(values["captured_at"].replace("Z", "+00:00"))).days
    return {"status": status, "reasons": reasons, "capture_age_days": capture_age,
            "warning": "Snapshot trend context, not a buy instruction, bottom forecast or loss guarantee. "
                       + ("User confirms completed bar." if values["bar_complete"] else "Bar is incomplete; values can change.")}
